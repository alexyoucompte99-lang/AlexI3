#!/usr/bin/env python3
"""API Conversions Meta : renvoi cote serveur de l'evenement Schedule du nouveau funnel.

Pourquoi : la page merci envoie Schedule depuis le navigateur seulement. Meta le recoit
(Events Manager) mais n'en rattache qu'une partie aux pubs (cookies perdus sur iOS/Safari,
arrivees sans fbclid, reservations faites ailleurs que sur la LP). L'API Conversions
renvoie la meme reservation depuis le serveur avec e-mail + telephone + cookies _fbp/_fbc :
Meta dedoublonne avec le Schedule navigateur (meme event_id, fenetre 48 h) et rattache mieux.

Deux sources, dans l'ordre :
  1. onglet « CAPI Schedule » du Sheet closing (rempli par la page merci via le pont,
     what=capi_schedule) : event_id du navigateur, e-mail, tel, cookies, user agent ;
  2. les reservations du calendrier « Appel Diagnostic - Club » presentes dans le Sheet closing
     mais jamais passees par la page merci (embed iClosed d'une autre page, setter...) :
     envoyees avec e-mail + telephone + nom seulement, event_id derive (pas de doublon possible).

Usage : python capi_send.py investisseurs30.xlsx data.json ../.capi-sent.json
            [--dry-run] [--test-code TESTxxxx] [--start AAAA-MM-JJ] [--since-days 7]
Env   : META_CAPI_TOKEN (token API Conversions du pixel ; absent = ne fait rien, sortie 0)
        META_PIXEL (defaut 1625394888920444)

Le registre .capi-sent.json (commite par le workflow) ne contient que des identifiants
haches : jamais d'e-mail ni de telephone dans le repo public.
"""
import datetime as dt
import hashlib
import json
import os
import re
import sys
import urllib.error
import urllib.request
from zoneinfo import ZoneInfo

PIXEL = os.environ.get("META_PIXEL", "1625394888920444")
TOKEN = os.environ.get("META_CAPI_TOKEN", "").strip()
API = "https://graph.facebook.com/v21.0"
PARIS = ZoneInfo("Europe/Paris")
NEW_CAL = re.compile(r"appel diagnostic\s*-\s*club", re.I)
# Premier jour pris en compte pour les reservations du Sheet sans passage par la page merci :
# avant cette date, le Schedule navigateur a deja plus de 48 h, le renvoyer ferait un doublon.
START_DEFAULT = "2026-09-15"
MERCI_URL = "https://thomas-mayol.com/strategie-privee/merci/"
RELAY_UA = "I3-Console-Relay/1.0 (+reservation iClosed relue dans le Sheet closing)"


def sha(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def norm_email(e):
    e = (e or "").strip().lower()
    return e if re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", e) else ""


def norm_phone(p):
    if isinstance(p, float) and p.is_integer():  # export xlsx Google « 33612345678.0 »
        p = int(p)
    d = re.sub(r"\D", "", str(p or ""))
    if re.match(r"^0\d{9}$", d):
        d = "33" + d[1:]
    return d if 8 <= len(d) <= 15 else ""


def norm_name(n):
    n = re.sub(r"\s+", " ", (n or "").strip().lower())
    if not n:
        return "", ""
    parts = n.split(" ")
    return parts[0], " ".join(parts[1:])


def user_data(email, phone, first, last, fbp="", fbc="", ua=""):
    ud = {}
    if email:
        ud["em"] = [sha(email)]
    if phone:
        ud["ph"] = [sha(phone)]
    if first:
        ud["fn"] = [sha(first)]
    if last:
        ud["ln"] = [sha(last)]
    if fbp:
        ud["fbp"] = fbp
    if fbc:
        ud["fbc"] = fbc
    if ua:
        ud["client_user_agent"] = ua
    return ud


def read_queue(xlsx_path):
    """Onglet « CAPI Schedule » -> liste de dicts (par en-tete)."""
    try:
        import openpyxl
    except ImportError:
        print("capi : openpyxl absent, file d'attente ignoree", file=sys.stderr)
        return []
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as e:
        print(f"capi : xlsx illisible ({e})", file=sys.stderr)
        return []
    if "CAPI Schedule" not in wb.sheetnames:
        return []
    ws = wb["CAPI Schedule"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    out = []
    for r in rows:
        if not r or not any(v not in (None, "") for v in r):
            continue
        d = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header))}
        out.append(d)
    return out


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = sys.argv[1:]
    if len(args) < 3:
        print(__doc__)
        sys.exit(1)
    xlsx_path, data_path, ledger_path = args[:3]
    dry = "--dry-run" in flags
    test_code = flags[flags.index("--test-code") + 1] if "--test-code" in flags else ""
    start = flags[flags.index("--start") + 1] if "--start" in flags else START_DEFAULT
    since_days = int(flags[flags.index("--since-days") + 1]) if "--since-days" in flags else 7

    if not TOKEN and not dry:
        print("capi : META_CAPI_TOKEN absent, rien envoye (ajouter le secret pour activer)")
        return

    now = dt.datetime.now(tz=dt.timezone.utc)
    now_ts = int(now.timestamp())
    min_ts = now_ts - since_days * 86400

    try:
        ledger = json.load(open(ledger_path, encoding="utf-8"))
    except Exception:
        ledger = {}
    ledger.setdefault("sent", {})   # event_id -> date d'envoi
    ledger.setdefault("rows", {})   # cle hachee (email|booking_date) -> event_id

    data = json.load(open(data_path, encoding="utf-8"))
    calls = [c for c in data.get("calls", []) if NEW_CAL.search(c.get("source") or "")]
    by_email = {}
    for c in calls:
        e = norm_email(c.get("mail"))
        if e:
            by_email.setdefault(e, []).append(c)
    # e-mails deja passes par la page merci (clic WhatsApp logue depuis la page) : le Schedule
    # navigateur a tres probablement ete envoye, on ne fabrique pas de second evenement.
    page_emails = {norm_email(w.get("email")) for w in data.get("wa_confirms", []) if (w.get("source") or "") == "page"}

    events, tags = [], []

    # --- 1. file d'attente de la page merci -------------------------------------------
    queue = read_queue(xlsx_path)
    queue_emails = set()
    for q in queue:
        eid = str(q.get("event id") or "").strip()
        if not re.match(r"^sch_[a-z0-9]{1,20}$", eid) or eid.startswith("sch_test"):
            continue
        email = norm_email(q.get("email"))
        if email:
            queue_emails.add(email)
        if eid in ledger["sent"]:
            continue
        # sans e-mail ni cookie, Meta n'aurait rien a faire correspondre : on ignore la ligne
        if not email and not str(q.get("fbp") or "").strip() and not str(q.get("fbc") or "").strip():
            continue
        try:
            ts = int(float(q.get("event time") or 0))
        except (TypeError, ValueError):
            ts = 0
        if not ts:
            d = q.get("date")
            if isinstance(d, dt.datetime):
                ts = int(d.replace(tzinfo=PARIS).timestamp())
        if ts < min_ts or ts > now_ts + 300:
            continue
        ts = min(ts, now_ts - 5)
        phone = norm_phone(q.get("tel"))
        first, last = norm_name(q.get("nom"))
        if email and (not phone or not first):
            for c in by_email.get(email, []):
                phone = phone or norm_phone(c.get("phone"))
                if not first:
                    first, last = norm_name(c.get("prospect"))
        ud = user_data(email, phone, first, last, str(q.get("fbp") or "").strip(),
                       str(q.get("fbc") or "").strip(), str(q.get("ua") or "").strip()[:300])
        if not ud.get("client_user_agent"):
            ud["client_user_agent"] = RELAY_UA
        events.append({
            "event_name": "Schedule", "event_time": ts, "event_id": eid,
            "action_source": "website",
            "event_source_url": str(q.get("url") or MERCI_URL).split("?")[0][:300] or MERCI_URL,
            "user_data": ud,
        })
        tags.append(("queue", eid, None))

    # --- 2. reservations du Sheet jamais passees par la page merci ----------------------
    for c in calls:
        bd = (c.get("booking_date") or "")[:10]
        email = norm_email(c.get("mail"))
        if not bd or not email or bd < start:
            continue
        if email in queue_emails or email in page_emails:
            continue
        key = "row:" + sha(email + "|" + bd)[:32]
        if key in ledger["rows"]:
            continue
        try:
            noon = dt.datetime.strptime(bd, "%Y-%m-%d").replace(hour=12, tzinfo=PARIS)
        except ValueError:
            continue
        ts = min(int(noon.timestamp()), now_ts - 5)
        if ts < min_ts:
            continue
        eid = "schs_" + sha(email + "|" + bd)[:16]
        first, last = norm_name(c.get("prospect"))
        ud = user_data(email, norm_phone(c.get("phone")), first, last, ua=RELAY_UA)
        events.append({
            "event_name": "Schedule", "event_time": ts, "event_id": eid,
            "action_source": "website", "event_source_url": MERCI_URL, "user_data": ud,
        })
        tags.append(("sheet", eid, key))

    if not events:
        print("capi : rien a envoyer")
        return

    if dry:
        print(f"capi (dry-run) : {len(events)} evenement(s) Schedule prets, pixel {PIXEL}")
        for ev, (src, eid, _) in zip(events, tags):
            ud = {k: (v[0][:10] + "…" if isinstance(v, list) else str(v)[:24]) for k, v in ev["user_data"].items()}
            when = dt.datetime.fromtimestamp(ev["event_time"], PARIS).strftime("%d/%m %H:%M")
            print(f"  [{src}] {eid} {when} {ud}")
        return

    payload = {"data": events, "access_token": TOKEN}
    if test_code:
        payload["test_event_code"] = test_code
    req = urllib.request.Request(f"{API}/{PIXEL}/events", data=json.dumps(payload).encode("utf-8"),
                                 headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            resp = json.load(r)
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", "replace")[:600]
        print(f"capi : refus Meta ({e.code}) : {body}", file=sys.stderr)
        return
    except Exception as e:
        print(f"capi : envoi impossible ({e})", file=sys.stderr)
        return
    received = resp.get("events_received")
    print(f"capi : {received} evenement(s) recu(s) par Meta"
          + (f" (test {test_code})" if test_code else ""))
    if not received:
        return
    stamp = dt.datetime.now(PARIS).strftime("%Y-%m-%d %H:%M")
    for src, eid, key in tags:
        ledger["sent"][eid] = stamp
        if key:
            ledger["rows"][key] = eid
    # on ne garde que 60 jours d'historique
    cutoff = (dt.datetime.now(PARIS) - dt.timedelta(days=60)).strftime("%Y-%m-%d")
    ledger["sent"] = {k: v for k, v in ledger["sent"].items() if v[:10] >= cutoff}
    keep_ids = set(ledger["sent"])
    ledger["rows"] = {k: v for k, v in ledger["rows"].items() if v in keep_ids}
    if not test_code:
        json.dump(ledger, open(ledger_path, "w", encoding="utf-8"), indent=0, sort_keys=True)


if __name__ == "__main__":
    main()
